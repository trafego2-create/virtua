"""Teste de ponta a ponta contra o banco da DATABASE_URL (use o Postgres local, nunca produção).

Recarrega a carteira no fim (scripts/carregar.py --forcar) para desfazer o que o teste alterou.
"""
import subprocess
import sys
from pathlib import Path

import psycopg
from fastapi.testclient import TestClient

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from app import config, db  # noqa: E402
from app.main import app  # noqa: E402

assert "localhost" in config.DATABASE_URL, "teste só roda no banco local"
H = {"x-virtua": "1"}


def novo_cliente():
    return TestClient(app, base_url="http://localhost")


def checar(cond, msg):
    print(("ok   " if cond else "FALHOU ") + msg)
    if not cond:
        raise SystemExit(1)


def main():
    with TestClient(app) as _:
        with db.pool.connection() as con:
            con.execute("delete from virtua.usuarios")
            con.execute("delete from virtua.auditoria where acao <> 'carga'")
        adm = novo_cliente()
        checar(adm.get("/api/me").json() == {"usuario": None, "primeiro_acesso": True}, "sem usuários pede primeiro acesso")
        r = adm.post("/api/primeiro-acesso", json={"nome": "Adm Teste", "email": "adm@teste.com", "senha": "12345678"}, headers=H)
        checar(r.status_code == 204, "primeiro acesso cria administrador e loga")
        checar(adm.post("/api/primeiro-acesso", json={"email": "x@x.com", "senha": "12345678"}, headers=H).status_code == 403,
               "primeiro acesso não funciona de novo")
        checar(adm.get("/api/me").json()["usuario"]["perfil"] == "administrador", "sessão do administrador")
        checar(adm.post("/api/usuarios", json={"nome": "Sem header", "email": "a@a.com", "perfil": "consulta", "senha": "12345678"}).status_code == 403,
               "escrita sem cabeçalho x-virtua é recusada")

        res = adm.get("/api/resumo").json()
        checar(res["carteira"] == {"ativa_em_dia": 260, "ativa_em_atraso": 1024, "inadimplencia_antiga": 3722}, f"carteira {res['carteira']}")
        checar(round(float(res["recebido"]), 2) == 1407796.87, "recebido total bate com o Virtua")

        lst = adm.get("/api/clientes", params={"carteira": "ativa", "ordem": "vencido"}).json()
        checar(lst["total"] == 1284 and len(lst["clientes"]) == 50, "lista da carteira ativa (1.284, 50 por página)")
        maior = lst["clientes"][0]
        cod = maior["codigo"]
        cpf = "".join(ch for ch in maior["cpf_cnpj"] if ch.isdigit())
        checar(adm.get("/api/clientes", params={"q": cpf[:6]}).json()["total"] >= 1, "busca por CPF")
        checar(adm.get("/api/clientes", params={"q": "HP"}).json()["total"] > 0, "busca por contrato")

        f = adm.get(f"/api/clientes/{cod}").json()
        checar(f["cliente"]["codigo"] == cod and f["parcelas"], "ficha do cliente com parcelas")
        venc = next(p for p in f["parcelas"] if p["status"] == "vencida")
        checar(float(venc["valor_atualizado"]) > float(venc["valor_previsto"]), "vencida mostra valor com juros e multa")

        checar(adm.patch(f"/api/clientes/{cod}", json={"telefone": "11999990000", "observacao": "teste"}, headers=H).status_code == 200, "edita contato")
        checar(adm.patch(f"/api/clientes/{cod}", json={"email": "inválido"}, headers=H).status_code == 400, "recusa e-mail inválido")

        r = adm.post(f"/api/clientes/{cod}/contratos", json={"produto": "Teste", "valor_total": "100", "qtde_parcelas": 3,
                                                            "primeiro_vencimento": "2027-01-31"}, headers=H)
        checar(r.status_code == 200, "cria contrato")
        conta_teste = r.json()["conta"]
        novas = [p for p in adm.get(f"/api/clientes/{cod}").json()["parcelas"] if p["conta"] == r.json()["conta"]]
        checar([p["vencimento"] for p in novas] == ["2027-01-31", "2027-02-28", "2027-03-31"], "vencimentos mensais respeitam fim de mês")
        checar(sum(float(p["valor_previsto"]) for p in novas) == 100.0 and float(novas[-1]["valor_previsto"]) == 33.34, "centavos vão para a última parcela")

        r = adm.patch(f"/api/parcelas/{venc['id']}", json={"acao": "baixar", "pago_em": "2026-10-20", "valor_pago": "300"}, headers=H)
        checar(r.status_code == 200, "dá baixa em parcela vencida")
        checar(adm.patch(f"/api/parcelas/{venc['id']}", json={"acao": "baixar", "pago_em": "2026-10-20", "valor_pago": "1"}, headers=H).status_code == 409,
               "não dá baixa duas vezes")
        checar(adm.patch(f"/api/parcelas/{venc['id']}", json={"acao": "estornar"}, headers=H).status_code == 200, "estorna")
        checar(adm.patch(f"/api/parcelas/{novas[0]['id']}", json={"acao": "baixar", "pago_em": "2099-01-01", "valor_pago": "1"}, headers=H).status_code == 400,
               "recusa pagamento no futuro")

        hist = adm.get(f"/api/clientes/{cod}").json()["historico"]
        checar(any(h["acao"] == "alterou" and h["tabela"] == "parcelas" and h["usuario"] == "adm@teste.com" for h in hist),
               "histórico registra quem deu baixa")
        checar(any(h["tabela"] == "clientes" and h["depois"]["telefone"] == "11999990000" for h in hist), "histórico registra edição do contato")

        checar(adm.put("/api/config/corte_carteira_ativa", json={"valor": "2025-01-01"}, headers=H).status_code == 200, "muda data de corte")
        checar(adm.get("/api/resumo").json()["carteira"].get("inadimplencia_antiga", 0) < 3722, "corte novo recalcula a carteira na hora")
        adm.put("/api/config/corte_carteira_ativa", json={"valor": "2026-01-01"}, headers=H)
        checar(adm.put("/api/config/multa_pct", json={"valor": "abc"}, headers=H).status_code == 400, "recusa percentual inválido")

        for email, perfil in [("fin@teste.com", "financeiro"), ("con@teste.com", "consulta"), ("ges@teste.com", "gestor")]:
            checar(adm.post("/api/usuarios", json={"nome": perfil, "email": email, "perfil": perfil, "senha": "12345678"}, headers=H).status_code == 200,
                   f"cria usuário {perfil}")
        sess = {}
        for email in ("fin@teste.com", "con@teste.com", "ges@teste.com"):
            sess[email] = novo_cliente()
            checar(sess[email].post("/api/login", json={"email": email, "senha": "12345678"}).status_code == 204, f"login {email}")
        con_, fin, ges = sess["con@teste.com"], sess["fin@teste.com"], sess["ges@teste.com"]
        checar(con_.get(f"/api/clientes/{cod}").status_code == 200, "consulta vê a ficha")
        checar(con_.get(f"/api/clientes/{cod}").json()["historico"] == [], "consulta não vê o histórico")
        checar(con_.patch(f"/api/clientes/{cod}", json={"uf": "SP"}, headers=H).status_code == 403, "consulta não edita")
        checar(ges.patch(f"/api/parcelas/{venc['id']}", json={"acao": "cancelar"}, headers=H).status_code == 403, "gestor não edita parcela")
        checar(ges.get("/api/auditoria").status_code == 200, "gestor vê histórico geral")
        checar(fin.patch(f"/api/clientes/{cod}", json={"uf": "SP"}, headers=H).status_code == 200, "financeiro edita")
        checar(fin.get("/api/usuarios").status_code == 403, "financeiro não gerencia usuários")
        checar(fin.put("/api/config/multa_pct", json={"valor": "3"}, headers=H).status_code == 403, "financeiro não muda configuração")

        uid = next(u["id"] for u in adm.get("/api/usuarios").json() if u["email"] == "fin@teste.com")
        adm.patch(f"/api/usuarios/{uid}", json={"ativo": False}, headers=H)
        checar(fin.get("/api/me").json()["usuario"] is None, "usuário desativado perde a sessão na hora")
        eu = next(u["id"] for u in adm.get("/api/usuarios").json() if u["email"] == "adm@teste.com")
        checar(adm.patch(f"/api/usuarios/{eu}", json={"perfil": "consulta"}, headers=H).status_code == 400, "admin não tira o próprio acesso")

        # 2ª entrega: inadimplência, previsão, painel, relatórios
        inad = adm.get("/api/inadimplencia").json()
        soma_faixas = round(sum(float(f["valor"]) for f in inad["faixas"].values()), 2)
        checar(soma_faixas == round(float(inad["valor_total"]), 2), "faixas de atraso somam o total da inadimplência")
        checar(set(inad["faixas"]) <= {"hoje", "1-7", "8-30", "31-60", "60+"}, "faixas são as do documento")
        so60 = adm.get("/api/inadimplencia", params={"faixa": "60+"}).json()
        checar(all(c["maior_atraso"] > 60 for c in so60["clientes"]), "filtro de faixa +60 só traz atraso acima de 60 dias")
        checar(all(float(c["valor_vencido"]) >= 1000 for c in adm.get("/api/inadimplencia", params={"minimo": "1000"}).json()["clientes"]),
               "filtro de valor mínimo")
        checar(adm.get("/api/inadimplencia", params={"minimo": "abc"}).status_code == 400, "valor mínimo inválido é recusado")
        prev = adm.get("/api/previsao", params={"dias": 60}).json()
        checar(len(prev["serie"]) == 61 and float(prev["em_7_dias"]) <= float(prev["em_30_dias"]) <= float(prev["futuro_total"]),
               "previsão: série diária e marcos crescentes")
        pn = adm.get("/api/painel", params={"de": "2025-01-01", "ate": "2025-12-31"}).json()
        checar(len(pn["meses"]) == 12 and round(sum(float(m["previsto"]) for m in pn["meses"]), 2) == round(float(pn["previsto"]), 2),
               "painel: 12 meses somam o previsto do período")
        checar(adm.get("/api/painel", params={"de": "2026-05-01", "ate": "2026-01-01"}).status_code == 400, "painel recusa período invertido")
        import io
        from openpyxl import load_workbook
        for tipo in ("clientes", "parcelas", "recebimentos", "inadimplencia"):
            x = con_.get(f"/api/relatorios/{tipo}", params={"carteira": "ativa"})
            ws = load_workbook(io.BytesIO(x.content)).active
            checar(x.status_code == 200 and ws.max_row > 1, f"relatório {tipo} em Excel abre ({ws.max_row - 1} linhas)")
        csv_txt = con_.get("/api/relatorios/clientes", params={"formato": "csv", "carteira": "ativa"}).content.decode("utf-8-sig")
        checar(csv_txt.splitlines()[0].startswith("Código;Nome") and len(csv_txt.splitlines()) == 1285, "relatório em CSV (1.284 clientes ativos)")
        checar(adm.get("/api/relatorios/xyz").status_code == 404, "relatório desconhecido dá 404")

        # importar do Virtua
        base_csv = (Path.home() / "Downloads" / "Base Virtua.csv").read_bytes()
        email_csv = (Path.home() / "Downloads" / "Grade Excel 04.csv").read_bytes()
        arqs = lambda b: [("arquivos", ("Base Virtua.csv", b, "text/csv")), ("arquivos", ("Grade Excel 04.csv", email_csv, "text/csv"))]  # noqa: E731
        pv = adm.post("/api/importacao/previa", files=arqs(base_csv), headers=H).json()
        checar(pv["pode_confirmar"] and pv["parcelas_novas"] == 0 and pv["qtd_pagamentos_novos"] == 0,
               "prévia com o mesmo relatório: confere e não muda nada")
        checar(len(pv["conflitos"]) >= 2 and pv["mantidos"]["contratos"] >= 1,
               "prévia avisa as alterações feitas aqui (baixa/estorno) e conta o contrato criado aqui")
        # relatório "de amanhã": uma parcela em aberto passa a paga, e a linha de totais acompanha
        linhas = base_csv.decode("latin1").split("\r\n")
        cab = linhas[0].split(";")
        iv, ip, ir = cab.index("Vlr. Previsto Parcela"), cab.index("Dta Pgto Parcela"), cab.index("Vlr. Realizado Parcela")
        alvo = next(i for i, l in enumerate(linhas[1:-2], start=1)
                    if (c := l.split(";"))[ip] == "" and c[iv] == "245" and c[cab.index("Nr. Acordo")] in ("", "0"))
        c = linhas[alvo].split(";"); c[ip], c[ir] = "20/10/2026", "245"; linhas[alvo] = ";".join(c)
        tot_i = len(linhas) - 2 if linhas[-1] == "" else len(linhas) - 1
        t = linhas[tot_i].split(";"); t[ir] = "1.408.041,87"; linhas[tot_i] = ";".join(t)
        amanha = "\r\n".join(linhas).encode("latin1")
        pv2 = adm.post("/api/importacao/previa", files=arqs(amanha), headers=H).json()
        checar(pv2["pode_confirmar"] and pv2["qtd_pagamentos_novos"] == 1 and float(pv2["valor_pagamentos_novos"]) == 245,
               "prévia mostra o pagamento novo do Virtua")
        checar(con_.post("/api/importacao/previa", files=arqs(amanha), headers=H).status_code == 403, "só administrador importa")
        r = adm.post("/api/importacao/confirmar", json={"token": pv2["token"]}, headers=H)
        checar(r.status_code == 200 and r.json()["recebido"] == "1408041.87", "importação confirmada com o total novo")
        checar(adm.post("/api/importacao/confirmar", json={"token": pv2["token"]}, headers=H).status_code == 410, "prévia não é usada duas vezes")
        f2 = adm.get(f"/api/clientes/{cod}").json()
        checar(any(p["conta"] == conta_teste for p in f2["parcelas"]) and f2["cliente"]["observacao"] == "teste",
               "contrato e observação criados aqui foram mantidos")
        checar(adm.get("/api/importacao").json()[0]["usuario"] == "adm@teste.com", "importação fica no histórico com quem fez")
        t[ir] = "9.999.999,99"; linhas[tot_i] = ";".join(t)
        pv3 = adm.post("/api/importacao/previa", files=arqs("\r\n".join(linhas).encode("latin1")), headers=H).json()
        checar(not pv3["pode_confirmar"], "total que não bate bloqueia a importação")
        checar(adm.post("/api/importacao/confirmar", json={"token": pv3["token"]}, headers=H).status_code == 409, "confirmar com total errado é recusado")
        x = adm.post("/api/importacao/previa", files=[("arquivos", ("foto.csv", b"a;b\n1;2\n3;4\n", "text/csv"))], headers=H)
        checar(x.status_code == 400 and "foto.csv" in x.json()["detail"], "arquivo errado é recusado com o nome dele")

        ruim = novo_cliente()
        for _ in range(8):
            ruim.post("/api/login", json={"email": "adm@teste.com", "senha": "errada"})
        checar(ruim.post("/api/login", json={"email": "adm@teste.com", "senha": "12345678"}).status_code == 429, "bloqueia após 8 tentativas erradas")
        checar(adm.get("/").status_code == 200 and "Cobranças" in adm.get("/clientes").text, "site responde")

    print("\nrecarregando a carteira para desfazer o teste…")
    subprocess.run([sys.executable, str(RAIZ / "scripts" / "carregar.py"), "--forcar"], check=True)
    with psycopg.connect(config.DATABASE_URL) as con:  # o pool já fechou junto com o TestClient
        # a carga preserva o que foi criado no sistema: apagar o que o teste criou
        con.execute("delete from virtua.parcelas where origem = 'sistema'")
        con.execute("delete from virtua.contratos where origem = 'sistema'")
        con.execute("delete from virtua.clientes where codigo like 'N-%'")
        con.execute("update virtua.clientes set observacao = null where observacao = 'teste'")
        con.execute("delete from virtua.usuarios")
        con.execute("delete from virtua.auditoria where acao <> 'carga'")
    print("tudo certo")


if __name__ == "__main__":
    main()
