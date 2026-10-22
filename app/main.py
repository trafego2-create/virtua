"""Virtua — API e site do sistema de cobranças."""
import re
import time
import uuid
from collections import defaultdict
from contextlib import asynccontextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from . import config, db, importacao, security

COOKIE = "virtua_sessao"
FUSO = ZoneInfo("America/Sao_Paulo")  # o container roda em UTC


def hoje() -> date:
    return datetime.now(FUSO).date()

PERFIS = ("administrador", "gestor", "financeiro", "consulta")
EDITA_CARTEIRA = {"administrador", "financeiro"}
VE_HISTORICO = {"administrador", "gestor"}


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.pool.open()
    yield
    db.pool.close()


app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None)


# ---------- sessão ----------
def usuario_atual(req: Request) -> dict | None:
    s = security.unsign(req.cookies.get(COOKIE))
    if not s:
        return None
    u = db.one("select * from virtua.usuarios where id = %s and ativo", (s["uid"],))
    if not u or security.password_version(u["senha_hash"]) != s.get("pv"):
        return None
    return u


def exigir(req: Request, perfis: set[str] | None = None) -> dict:
    u = usuario_atual(req)
    if not u:
        raise HTTPException(401, "Faça login")
    if perfis and u["perfil"] not in perfis:
        raise HTTPException(403, "Seu perfil não permite esta ação")
    if req.method not in ("GET", "HEAD") and req.headers.get("x-virtua") != "1":
        raise HTTPException(403, "Requisição recusada")
    return u


def publico(u: dict) -> dict:
    return {"id": u["id"], "nome": u["nome"], "email": u["email"], "perfil": u["perfil"]}


def gravar_sessao(resp: Response, u: dict, lembrar: bool) -> None:
    ttl = 30 * 86400 if lembrar else 12 * 3600
    token = security.sign({"uid": u["id"], "pv": security.password_version(u["senha_hash"])}, ttl)
    resp.set_cookie(COOKIE, token, max_age=ttl if lembrar else None, httponly=True,
                    secure=config.COOKIE_SECURE, samesite="strict")


_tentativas: dict[str, list[float]] = defaultdict(list)


def muitas_tentativas(chave: str) -> bool:
    agora = time.time()
    _tentativas[chave] = [t for t in _tentativas[chave] if agora - t < 900]
    return len(_tentativas[chave]) >= 8


# ---------- validação ----------
def dinheiro(v, campo: str) -> Decimal:
    try:
        d = Decimal(str(v).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError):
        raise HTTPException(400, f"{campo}: valor inválido")
    if d < 0:
        raise HTTPException(400, f"{campo}: não pode ser negativo")
    return d


def dia(v, campo: str) -> date:
    try:
        return date.fromisoformat(str(v))
    except ValueError:
        raise HTTPException(400, f"{campo}: data inválida")


def texto(v, maximo=200) -> str | None:
    v = (v or "").strip() if isinstance(v, str) or v is None else str(v).strip()
    return v[:maximo] or None


async def corpo(req: Request) -> dict:
    try:
        b = await req.json()
    except ValueError:
        raise HTTPException(400, "JSON inválido")
    if not isinstance(b, dict):
        raise HTTPException(400, "JSON inválido")
    return b


# ---------- login ----------
@app.get("/healthz")
def healthz():
    db.one("select 1")
    return {"ok": True}


@app.get("/api/me")
def me(req: Request):
    u = usuario_atual(req)
    if not u:
        primeiro = db.one("select count(*) as n from virtua.usuarios")["n"] == 0
        return {"usuario": None, "primeiro_acesso": primeiro}
    return {"usuario": publico(u)}


@app.post("/api/login")
async def login(req: Request):
    b = await corpo(req)
    email = (b.get("email") or "").strip().lower()
    ip = req.client.host if req.client else "?"
    if muitas_tentativas(ip) or muitas_tentativas(email):
        raise HTTPException(429, "Muitas tentativas. Espere 15 minutos.")
    u = db.one("select * from virtua.usuarios where email = %s and ativo", (email,))
    if not u or not security.verify_password(b.get("senha") or "", u["senha_hash"]):
        _tentativas[ip].append(time.time())
        _tentativas[email].append(time.time())
        raise HTTPException(401, "E-mail ou senha incorretos")
    resp = Response(status_code=204)
    gravar_sessao(resp, u, bool(b.get("lembrar")))
    return resp


@app.post("/api/logout")
def logout():
    resp = Response(status_code=204)
    resp.delete_cookie(COOKIE)
    return resp


@app.post("/api/primeiro-acesso")
async def primeiro_acesso(req: Request):
    """Cria o primeiro administrador. Só funciona com a tabela de usuários vazia."""
    if req.headers.get("x-virtua") != "1":
        raise HTTPException(403, "Requisição recusada")
    b = await corpo(req)
    senha = b.get("senha") or ""
    if len(senha) < 8:
        raise HTTPException(400, "A senha precisa ter pelo menos 8 caracteres")
    with db.tx("primeiro acesso") as con:
        con.execute("lock table virtua.usuarios in exclusive mode")
        if con.execute("select count(*) as n from virtua.usuarios").fetchone()["n"]:
            raise HTTPException(403, "O primeiro acesso já foi feito")
        u = con.execute(
            "insert into virtua.usuarios (nome, email, perfil, senha_hash) values (%s, %s, 'administrador', %s) returning *",
            (texto(b.get("nome")) or "Administrador", validar_email(b.get("email")), security.hash_password(senha)),
        ).fetchone()
    resp = Response(status_code=204)
    gravar_sessao(resp, u, False)
    return resp


@app.post("/api/senha")
async def trocar_senha(req: Request):
    u = exigir(req)
    b = await corpo(req)
    if not security.verify_password(b.get("atual") or "", u["senha_hash"]):
        raise HTTPException(400, "Senha atual incorreta")
    nova = b.get("nova") or ""
    if len(nova) < 8:
        raise HTTPException(400, "A senha nova precisa ter pelo menos 8 caracteres")
    with db.tx(u["email"]) as con:
        u = con.execute("update virtua.usuarios set senha_hash = %s where id = %s returning *",
                        (security.hash_password(nova), u["id"])).fetchone()
    resp = Response(status_code=204)
    gravar_sessao(resp, u, False)
    return resp


def validar_email(v) -> str:
    e = (v or "").strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", e):
        raise HTTPException(400, "E-mail inválido")
    return e


# ---------- resumo ----------
@app.get("/api/resumo")
def resumo(req: Request):
    exigir(req)
    p = db.one("""
        select coalesce(sum(valor_previsto) filter (where status in ('vencida','a_vencer')), 0) as a_receber,
               coalesce(sum(valor_pago), 0)                                                     as recebido,
               coalesce(sum(valor_previsto) filter (where status = 'vencida'), 0)               as vencido,
               coalesce(sum(valor_previsto) filter (where status = 'a_vencer'), 0)              as a_vencer,
               coalesce(sum(valor_previsto) filter (where status = 'vencida'
                        and vencimento >= date_trunc('year', virtua.hoje())), 0)                 as vencido_ano
        from virtua.v_parcelas""")
    c = db.all("select carteira, count(*) as n from virtua.v_clientes group by carteira")
    return {**p, "ano": hoje().year, "carteira": {r["carteira"]: r["n"] for r in c}}


# ---------- clientes ----------
ORDENS = {
    "nome": "nome",
    "vencido": "total_vencido desc, nome",
    "atraso": "maior_atraso desc, nome",
    "pagamento": "ultimo_pagamento desc nulls last, nome",
}


@app.get("/api/clientes")
def listar_clientes(req: Request, q: str = "", carteira: str = "", ordem: str = "nome", pagina: int = 1):
    exigir(req)
    onde, params = [], []
    if q.strip():
        termo = q.strip()
        digitos = re.sub(r"\D", "", termo)
        cond = "(nome ilike %s or email ilike %s or codigo = %s"
        params += [f"%{termo}%", f"%{termo}%", termo]
        if len(digitos) >= 3:
            cond += " or regexp_replace(cpf_cnpj, '\\D', '', 'g') like %s or telefone like %s"
            params += [f"%{digitos}%", f"%{digitos}%"]
        cond += " or exists (select 1 from virtua.contratos k where k.codigo_cliente = v.codigo and k.conta ilike %s))"
        params.append(f"%{termo}%")
        onde.append(cond)
    if carteira in ("ativa_em_dia", "ativa_em_atraso", "inadimplencia_antiga"):
        onde.append("carteira = %s")
        params.append(carteira)
    elif carteira == "ativa":
        onde.append("carteira like 'ativa%%'")
    where = ("where " + " and ".join(onde)) if onde else ""
    pagina = max(1, pagina)
    total = db.one(f"select count(*) as n from virtua.v_clientes v {where}", params)["n"]
    linhas = db.all(
        f"""select codigo, cpf_cnpj, nome, telefone, email, uf, carteira, total_em_aberto, total_vencido,
                   parcelas_vencidas, maior_atraso, ultimo_pagamento
            from virtua.v_clientes v {where} order by {ORDENS.get(ordem, 'nome')} limit 50 offset %s""",
        params + [(pagina - 1) * 50])
    return {"total": total, "pagina": pagina, "por_pagina": 50, "clientes": linhas}


@app.post("/api/clientes")
async def criar_cliente(req: Request):
    u = exigir(req, EDITA_CARTEIRA)
    b = await corpo(req)
    nome, doc = texto(b.get("nome")), texto(b.get("cpf_cnpj"), 20)
    if not nome or not doc:
        raise HTTPException(400, "Nome e CPF/CNPJ são obrigatórios")
    if db.one("select 1 from virtua.clientes where regexp_replace(cpf_cnpj, '\\D', '', 'g') = %s",
              (re.sub(r"\D", "", doc),)):
        raise HTTPException(409, "Já existe cliente com este CPF/CNPJ")
    with db.tx(u["email"]) as con:
        c = con.execute(
            """insert into virtua.clientes (codigo, cpf_cnpj, nome, telefone, email, uf)
               values ('N-' || nextval('virtua.clientes_novos_seq'), %s, %s, %s, %s, %s) returning codigo""",
            (doc, nome, texto(b.get("telefone"), 30),
             validar_email(b["email"]) if b.get("email") else None, texto(b.get("uf"), 2)),
        ).fetchone()
    return c


@app.get("/api/clientes/{codigo}")
def ficha(codigo: str, req: Request):
    u = exigir(req)
    c = db.one("select * from virtua.v_clientes where codigo = %s", (codigo,))
    if not c:
        raise HTTPException(404, "Cliente não encontrado")
    contratos = db.all(
        """select k.*, count(p.id) as parcelas, coalesce(sum(p.valor_pago), 0) as recebido
           from virtua.contratos k left join virtua.parcelas p on p.conta = k.conta and p.nr_acordo is null
           where k.codigo_cliente = %s group by k.conta order by k.criado_em, k.conta""", (codigo,))
    acordos = db.all("select * from virtua.acordos where codigo_cliente = %s order by data_acordo desc nulls last", (codigo,))
    parcelas = db.all("select * from virtua.v_parcelas where codigo_cliente = %s order by vencimento, id", (codigo,))
    historico = []
    if u["perfil"] in VE_HISTORICO:
        historico = db.all(
            """select em, usuario, tabela, registro, acao, antes, depois from virtua.auditoria
               where (tabela = 'clientes' and registro = %s)
                  or (tabela in ('parcelas','contratos','acordos')
                      and coalesce(depois->>'codigo_cliente', antes->>'codigo_cliente') = %s)
               order by em desc limit 200""", (codigo, codigo))
    return {"cliente": c, "contratos": contratos, "acordos": acordos, "parcelas": parcelas, "historico": historico}


@app.patch("/api/clientes/{codigo}")
async def editar_cliente(codigo: str, req: Request):
    u = exigir(req, EDITA_CARTEIRA)
    b = await corpo(req)
    campos = {}
    if "nome" in b:
        campos["nome"] = texto(b["nome"])
        if not campos["nome"]:
            raise HTTPException(400, "Nome não pode ficar vazio")
    if "telefone" in b:
        campos["telefone"] = texto(b["telefone"], 30)
    if "email" in b:
        campos["email"] = validar_email(b["email"]) if b["email"] else None
    if "uf" in b:
        campos["uf"] = (texto(b["uf"], 2) or "").upper() or None
    if "observacao" in b:
        campos["observacao"] = texto(b["observacao"], 2000)
    if not campos:
        raise HTTPException(400, "Nada para alterar")
    sets = ", ".join(f"{k} = %s" for k in campos)
    with db.tx(u["email"]) as con:
        r = con.execute(f"update virtua.clientes set {sets} where codigo = %s returning codigo",
                        list(campos.values()) + [codigo]).fetchone()
    if not r:
        raise HTTPException(404, "Cliente não encontrado")
    return {"ok": True}


# ---------- contratos e parcelas ----------
def somar_meses(d: date, n: int) -> date:
    m = d.month - 1 + n
    a, m = d.year + m // 12, m % 12 + 1
    dias = [31, 29 if a % 4 == 0 and (a % 100 or a % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return date(a, m, min(d.day, dias[m - 1]))


@app.post("/api/clientes/{codigo}/contratos")
async def criar_contrato(codigo: str, req: Request):
    u = exigir(req, EDITA_CARTEIRA)
    b = await corpo(req)
    total = dinheiro(b.get("valor_total"), "Valor total")
    try:
        n = int(b.get("qtde_parcelas"))
    except (TypeError, ValueError):
        raise HTTPException(400, "Quantidade de parcelas inválida")
    if not 1 <= n <= 120 or total <= 0:
        raise HTTPException(400, "Informe valor total maior que zero e de 1 a 120 parcelas")
    primeiro = dia(b.get("primeiro_vencimento"), "Primeiro vencimento")
    base = (total / n).quantize(Decimal("0.01"))
    valores = [base] * (n - 1) + [total - base * (n - 1)]  # diferença de centavos na última
    conta = "C-" + uuid.uuid4().hex[:10].upper()
    with db.tx(u["email"]) as con:
        if not con.execute("select 1 from virtua.clientes where codigo = %s", (codigo,)).fetchone():
            raise HTTPException(404, "Cliente não encontrado")
        con.execute("insert into virtua.contratos (conta, codigo_cliente, produto, origem) values (%s, %s, %s, 'sistema')",
                    (conta, codigo, texto(b.get("produto"))))
        for i, v in enumerate(valores):
            con.execute(
                """insert into virtua.parcelas (codigo_cliente, conta, nr_parcela, vencimento, valor_previsto, situacao, origem)
                   values (%s, %s, %s, %s, %s, 'aberta', 'sistema')""",
                (codigo, conta, i + 1, somar_meses(primeiro, i), v))
    return {"conta": conta}


@app.patch("/api/parcelas/{pid}")
async def alterar_parcela(pid: int, req: Request):
    u = exigir(req, EDITA_CARTEIRA)
    b = await corpo(req)
    acao = b.get("acao")
    with db.tx(u["email"]) as con:
        p = con.execute("select * from virtua.parcelas where id = %s for update", (pid,)).fetchone()
        if not p:
            raise HTTPException(404, "Parcela não encontrada")
        if acao == "baixar":
            if p["situacao"] != "aberta":
                raise HTTPException(409, "Só dá para dar baixa em parcela em aberto")
            pago_em = dia(b.get("pago_em"), "Data do pagamento")
            if pago_em > hoje():
                raise HTTPException(400, "Data do pagamento não pode ser no futuro")
            con.execute("update virtua.parcelas set situacao = 'paga', pago_em = %s, valor_pago = %s where id = %s",
                        (pago_em, dinheiro(b.get("valor_pago"), "Valor pago"), pid))
        elif acao == "estornar":
            if p["situacao"] != "paga":
                raise HTTPException(409, "Só dá para estornar parcela paga")
            con.execute("update virtua.parcelas set situacao = 'aberta', pago_em = null, valor_pago = 0 where id = %s", (pid,))
        elif acao == "cancelar":
            if p["situacao"] != "aberta":
                raise HTTPException(409, "Só dá para cancelar parcela em aberto")
            con.execute("update virtua.parcelas set situacao = 'cancelada' where id = %s", (pid,))
        elif acao == "reabrir":
            if p["situacao"] != "cancelada":
                raise HTTPException(409, "Só dá para reabrir parcela cancelada")
            con.execute("update virtua.parcelas set situacao = 'aberta' where id = %s", (pid,))
        elif acao == "editar":
            if p["situacao"] != "aberta":
                raise HTTPException(409, "Só dá para editar parcela em aberto")
            con.execute("update virtua.parcelas set vencimento = %s, valor_previsto = %s where id = %s",
                        (dia(b.get("vencimento"), "Vencimento"), dinheiro(b.get("valor_previsto"), "Valor"), pid))
        else:
            raise HTTPException(400, "Ação inválida")
    return {"ok": True}


# ---------- configurações e usuários (administrador) ----------
@app.get("/api/config")
def ler_config(req: Request):
    exigir(req)
    return db.all("select chave, valor, descricao, atualizado_em from virtua.config order by chave")


@app.put("/api/config/{chave}")
async def gravar_config(chave: str, req: Request):
    u = exigir(req, {"administrador"})
    b = await corpo(req)
    valor = str(b.get("valor", "")).strip()
    if chave == "corte_carteira_ativa":
        valor = dia(valor, "Data de corte").isoformat()
    elif chave in ("multa_pct", "juros_mes_pct"):
        d = dinheiro(valor, "Percentual")
        if d > 100:
            raise HTTPException(400, "Percentual acima de 100")
        valor = str(d)
    else:
        raise HTTPException(404, "Configuração desconhecida")
    with db.tx(u["email"]) as con:
        con.execute("update virtua.config set valor = %s where chave = %s", (valor, chave))
    return {"ok": True}


@app.get("/api/usuarios")
def listar_usuarios(req: Request):
    exigir(req, {"administrador"})
    return db.all("select id, nome, email, perfil, ativo, criado_em from virtua.usuarios order by nome")


@app.post("/api/usuarios")
async def criar_usuario(req: Request):
    u = exigir(req, {"administrador"})
    b = await corpo(req)
    perfil, senha = b.get("perfil"), b.get("senha") or ""
    if perfil not in PERFIS:
        raise HTTPException(400, "Perfil inválido")
    if len(senha) < 8:
        raise HTTPException(400, "A senha precisa ter pelo menos 8 caracteres")
    email = validar_email(b.get("email"))
    if db.one("select 1 from virtua.usuarios where email = %s", (email,)):
        raise HTTPException(409, "Já existe usuário com este e-mail")
    with db.tx(u["email"]) as con:
        con.execute("insert into virtua.usuarios (nome, email, perfil, senha_hash) values (%s, %s, %s, %s)",
                    (texto(b.get("nome")) or email, email, perfil, security.hash_password(senha)))
    return {"ok": True}


@app.patch("/api/usuarios/{uid}")
async def alterar_usuario(uid: int, req: Request):
    u = exigir(req, {"administrador"})
    b = await corpo(req)
    campos = {}
    if "perfil" in b:
        if b["perfil"] not in PERFIS:
            raise HTTPException(400, "Perfil inválido")
        campos["perfil"] = b["perfil"]
    if "ativo" in b:
        campos["ativo"] = bool(b["ativo"])
    if b.get("senha"):
        if len(b["senha"]) < 8:
            raise HTTPException(400, "A senha precisa ter pelo menos 8 caracteres")
        campos["senha_hash"] = security.hash_password(b["senha"])
    if not campos:
        raise HTTPException(400, "Nada para alterar")
    if uid == u["id"] and (campos.get("perfil", "administrador") != "administrador" or campos.get("ativo") is False):
        raise HTTPException(400, "Você não pode tirar o próprio acesso de administrador")
    sets = ", ".join(f"{k} = %s" for k in campos)
    with db.tx(u["email"]) as con:
        if not con.execute(f"update virtua.usuarios set {sets} where id = %s returning id",
                           list(campos.values()) + [uid]).fetchone():
            raise HTTPException(404, "Usuário não encontrado")
    return {"ok": True}


@app.get("/api/auditoria")
def auditoria(req: Request, pagina: int = 1):
    exigir(req, VE_HISTORICO)
    return db.all("select * from virtua.auditoria order by em desc limit 100 offset %s", ((max(1, pagina) - 1) * 100,))


# ---------- gestão: inadimplência, previsão, painel ----------
FAIXAS = ("hoje", "1-7", "8-30", "31-60", "60+")
SQL_FAIXA = """case when virtua.hoje() - p.vencimento = 0 then 'hoje'
                    when virtua.hoje() - p.vencimento <= 7 then '1-7'
                    when virtua.hoje() - p.vencimento <= 30 then '8-30'
                    when virtua.hoje() - p.vencimento <= 60 then '31-60'
                    else '60+' end"""


def data_opcional(v: str | None, campo: str) -> date | None:
    return dia(v, campo) if v else None


def filtro_cliente(q: str, carteira: str, alias: str = "c") -> tuple[str, dict]:
    """Condições sobre v_clientes (busca + carteira), com parâmetros nomeados."""
    conds, params = [], {}
    if q.strip():
        termo = q.strip()
        params |= {"q_like": f"%{termo}%", "q_exato": termo}
        cond = f"({alias}.nome ilike %(q_like)s or {alias}.email ilike %(q_like)s or {alias}.codigo = %(q_exato)s"
        digitos = re.sub(r"\D", "", termo)
        if len(digitos) >= 3:
            params["q_dig"] = f"%{digitos}%"
            cond += f" or regexp_replace({alias}.cpf_cnpj, '\\D', '', 'g') like %(q_dig)s or {alias}.telefone like %(q_dig)s"
        conds.append(cond + ")")
    if carteira == "ativa":
        conds.append(f"{alias}.carteira like 'ativa%%'")
    elif carteira in ("ativa_em_dia", "ativa_em_atraso", "inadimplencia_antiga"):
        conds.append(f"{alias}.carteira = %(carteira)s")
        params["carteira"] = carteira
    return " and ".join(conds) or "true", params


def consulta_inadimplencia(q, carteira, faixa, de, ate, minimo, maximo, ordem, limite, offset):
    cond_cli, params = filtro_cliente(q, carteira)
    params |= {"de": data_opcional(de, "De"), "ate": data_opcional(ate, "Até"),
               "faixa": faixa if faixa in FAIXAS else None,
               "min": dinheiro(minimo, "Valor mínimo") if minimo else None,
               "max": dinheiro(maximo, "Valor máximo") if maximo else None}
    base = f"""with f as (
          select p.codigo_cliente, p.vencimento, p.valor_previsto, p.valor_atualizado,
                 virtua.hoje() - p.vencimento as dias, {SQL_FAIXA} as faixa
          from virtua.v_parcelas p join virtua.v_clientes c on c.codigo = p.codigo_cliente
          where p.situacao = 'aberta' and p.vencimento <= virtua.hoje() and {cond_cli}
            and (%(de)s::date is null or p.vencimento >= %(de)s::date)
            and (%(ate)s::date is null or p.vencimento <= %(ate)s::date))"""
    resumo = db.all(base + """
        select faixa, count(*) as parcelas, count(distinct codigo_cliente) as clientes,
               sum(valor_previsto) as valor, sum(valor_atualizado) as valor_atualizado
        from f group by faixa""", params)
    ordens = {"valor": "valor_vencido desc", "atraso": "maior_atraso desc", "nome": "c.nome",
              "antigo": "vencimento_mais_antigo"}
    agrupado = base + """, g as (
          select codigo_cliente, count(*) as parcelas_vencidas, sum(valor_previsto) as valor_vencido,
                 sum(valor_atualizado) as valor_atualizado, max(dias) as maior_atraso,
                 min(vencimento) as vencimento_mais_antigo
          from f where %(faixa)s::text is null or faixa = %(faixa)s::text
          group by codigo_cliente
          having (%(min)s::numeric is null or sum(valor_previsto) >= %(min)s::numeric)
             and (%(max)s::numeric is null or sum(valor_previsto) <= %(max)s::numeric))"""
    total = db.one(agrupado + " select count(*) as n, coalesce(sum(valor_vencido), 0) as valor from g", params)
    linhas = db.all(agrupado + f"""
        select c.codigo, c.nome, c.cpf_cnpj, c.telefone, c.email, c.carteira, c.ultimo_pagamento, g.*
        from g join virtua.v_clientes c on c.codigo = g.codigo_cliente
        order by {ordens.get(ordem, 'valor_vencido desc')}, c.nome
        limit %(lim)s offset %(off)s""", params | {"lim": limite, "off": offset})
    return resumo, total, linhas


@app.get("/api/inadimplencia")
def inadimplencia(req: Request, q: str = "", carteira: str = "ativa", faixa: str = "", de: str = "", ate: str = "",
                  minimo: str = "", maximo: str = "", ordem: str = "valor", pagina: int = 1):
    exigir(req)
    pagina = max(1, pagina)
    resumo, total, linhas = consulta_inadimplencia(q, carteira, faixa, de, ate, minimo, maximo, ordem, 50, (pagina - 1) * 50)
    return {"faixas": {r["faixa"]: r for r in resumo}, "total": total["n"], "valor_total": total["valor"],
            "pagina": pagina, "por_pagina": 50, "clientes": linhas}


@app.get("/api/previsao")
def previsao(req: Request, dias: int = 30):
    exigir(req)
    dias = min(max(dias, 7), 365)
    marcos = db.one("""
        select coalesce(sum(valor_previsto) filter (where vencimento = virtua.hoje()), 0)          as hoje,
               coalesce(sum(valor_previsto) filter (where vencimento = virtua.hoje() + 1), 0)      as amanha,
               coalesce(sum(valor_previsto) filter (where vencimento <= virtua.hoje() + 7), 0)     as em_7_dias,
               coalesce(sum(valor_previsto) filter (where vencimento <= virtua.hoje() + 30), 0)    as em_30_dias,
               coalesce(sum(valor_previsto), 0)                                                    as futuro_total,
               count(*)                                                                            as parcelas_futuras
        from virtua.parcelas where situacao = 'aberta' and vencimento >= virtua.hoje()""")
    vencido = db.one("select coalesce(sum(valor_previsto), 0) as v from virtua.parcelas "
                     "where situacao = 'aberta' and vencimento < virtua.hoje()")["v"]
    serie = db.all("""
        select d::date as dia, coalesce(sum(p.valor_previsto), 0) as valor, count(p.id) as parcelas
        from generate_series(virtua.hoje(), virtua.hoje() + %s, interval '1 day') d
        left join virtua.parcelas p on p.vencimento = d::date and p.situacao = 'aberta'
        group by d order by d""", (dias,))
    parcelas = db.all("""
        select p.id, p.vencimento, p.valor_previsto, p.nr_parcela, p.nr_acordo, p.conta, c.codigo, c.nome
        from virtua.parcelas p join virtua.clientes c on c.codigo = p.codigo_cliente
        where p.situacao = 'aberta' and p.vencimento between virtua.hoje() and virtua.hoje() + %s
        order by p.vencimento, c.nome limit 300""", (dias,))
    return {**marcos, "vencido": vencido, "dias": dias, "serie": serie, "parcelas": parcelas}


@app.get("/api/painel")
def painel(req: Request, de: str = "", ate: str = ""):
    exigir(req)
    agora = hoje()
    d1 = data_opcional(de, "De") or date(agora.year, 1, 1)
    d2 = data_opcional(ate, "Até") or agora
    if d2 < d1:
        raise HTTPException(400, "A data final é anterior à inicial")
    p = {"de": d1, "ate": d2}
    tot = db.one("""
        select coalesce(sum(valor_previsto) filter (where vencimento between %(de)s and %(ate)s
                        and status not in ('renegociada','cancelada')), 0)                         as previsto,
               coalesce(sum(valor_pago) filter (where vencimento between %(de)s and %(ate)s), 0)   as realizado,
               coalesce(sum(valor_pago) filter (where pago_em between %(de)s and %(ate)s), 0)      as recebido,
               coalesce(sum(valor_previsto) filter (where vencimento between %(de)s and %(ate)s
                        and status = 'vencida'), 0)                                                 as vencido_periodo,
               coalesce(sum(valor_previsto) filter (where status in ('vencida','a_vencer')), 0)    as a_receber,
               coalesce(sum(valor_previsto) filter (where status = 'vencida'), 0)                  as vencido_total,
               coalesce(sum(valor_previsto) filter (where status = 'a_vencer'), 0)                 as a_vencer,
               count(*) filter (where nosso_numero is not null and vencimento between %(de)s and %(ate)s)                      as boletos_emitidos,
               count(*) filter (where nosso_numero is not null and vencimento between %(de)s and %(ate)s and status = 'paga')  as boletos_pagos,
               count(*) filter (where nosso_numero is not null and vencimento between %(de)s and %(ate)s and status = 'vencida') as boletos_vencidos,
               count(*) filter (where nosso_numero is not null and vencimento between %(de)s and %(ate)s and status = 'a_vencer') as boletos_em_aberto
        from virtua.v_parcelas""", p)
    meses = db.all("""
        with m as (select generate_series(date_trunc('month', %(de)s::date), date_trunc('month', %(ate)s::date),
                                          interval '1 month')::date as mes)
        select m.mes,
               coalesce((select sum(valor_previsto) from virtua.v_parcelas x where date_trunc('month', x.vencimento) = m.mes
                         and x.status not in ('renegociada','cancelada')), 0) as previsto,
               coalesce((select sum(valor_pago) from virtua.v_parcelas x where date_trunc('month', x.vencimento) = m.mes), 0) as realizado,
               coalesce((select sum(valor_pago) from virtua.v_parcelas x where date_trunc('month', x.pago_em) = m.mes), 0) as recebido
        from m order by m.mes""", p)
    return {"de": d1, "ate": d2, **tot, "meses": meses}


# ---------- importar do Virtua (administrador) ----------
# A prévia fica em memória até ser confirmada (1 worker; ver Dockerfile). Expira em 30 minutos.
_importacoes: dict[str, dict] = {}
MAX_ARQUIVO = 60 * 1024 * 1024


@app.get("/api/importacao")
def historico_importacoes(req: Request):
    exigir(req, {"administrador"})
    return db.all("""select em, usuario, depois from virtua.auditoria where acao = 'carga'
                     order by em desc limit 20""")


@app.post("/api/importacao/previa")
def previa_importacao(req: Request, arquivos: list[UploadFile] = File(...)):
    u = exigir(req, {"administrador"})
    if not 1 <= len(arquivos) <= 2:
        raise HTTPException(400, "Envie o relatório completo do Virtua e, se tiver, o relatório com e-mail")
    lidos = []
    for a in arquivos:
        conteudo = a.file.read(MAX_ARQUIVO + 1)
        if len(conteudo) > MAX_ARQUIVO:
            raise HTTPException(400, f"{a.filename}: arquivo maior que 60 MB")
        lidos.append((a.filename or "arquivo", conteudo))
    try:
        dados = importacao.ler_arquivos(lidos)
    except (importacao.ArquivoInvalido, ValueError, KeyError, InvalidOperation) as e:
        raise HTTPException(400, f"Não consegui ler: {e}")
    with db.pool.connection() as con:
        p = importacao.previa(con, dados)
        ultima = con.execute("select coalesce(max(id), 0) as id from virtua.auditoria where acao = 'carga'").fetchone()["id"]
    agora = time.time()
    for k in [k for k, v in _importacoes.items() if agora - v["criado"] > 1800]:
        del _importacoes[k]
    token = uuid.uuid4().hex
    _importacoes[token] = {"dados": dados, "usuario": u["id"], "criado": agora, "ultima_carga": ultima,
                           "pode": p["pode_confirmar"]}
    return {**p, "token": token}


@app.post("/api/importacao/confirmar")
async def confirmar_importacao(req: Request):
    u = exigir(req, {"administrador"})
    b = await corpo(req)
    imp = _importacoes.get(str(b.get("token")))
    if not imp or imp["usuario"] != u["id"] or time.time() - imp["criado"] > 1800:
        raise HTTPException(410, "A prévia expirou. Envie os arquivos de novo.")
    if not imp["pode"]:
        raise HTTPException(409, "A conferência com os totais do Virtua não bateu; a importação foi bloqueada")
    with db.tx(u["email"]) as con:
        con.execute("lock table virtua.parcelas in exclusive mode")
        ultima = con.execute("select coalesce(max(id), 0) as id from virtua.auditoria where acao = 'carga'").fetchone()["id"]
        if ultima != imp["ultima_carga"]:
            raise HTTPException(409, "Outra importação foi feita depois desta prévia. Envie os arquivos de novo.")
        resumo = importacao.aplicar(con, imp["dados"], u["email"])
    _importacoes.pop(str(b.get("token")), None)
    return resumo


# ---------- relatórios (Excel / CSV) ----------
RELATORIOS = {
    "clientes": "Clientes",
    "parcelas": "Cobranças",
    "recebimentos": "Recebimentos",
    "inadimplencia": "Inadimplência",
}


def dados_relatorio(tipo, q, carteira, status, faixa, de, ate, minimo, maximo):
    if tipo == "clientes":
        cond, params = filtro_cliente(q, carteira)
        linhas = db.all(f"""select codigo, nome, cpf_cnpj, telefone, email, uf, carteira, total_contratado, total_recebido,
                                   total_em_aberto, total_vencido, parcelas_vencidas, maior_atraso, ultimo_pagamento
                            from virtua.v_clientes c where {cond} order by nome""", params)
        cab = [("codigo", "Código"), ("nome", "Nome"), ("cpf_cnpj", "CPF/CNPJ"), ("telefone", "Telefone"), ("email", "E-mail"),
               ("uf", "UF"), ("carteira", "Carteira"), ("total_contratado", "Contratado"), ("total_recebido", "Recebido"),
               ("total_em_aberto", "Em aberto"), ("total_vencido", "Atrasado"), ("parcelas_vencidas", "Parcelas vencidas"),
               ("maior_atraso", "Maior atraso (dias)"), ("ultimo_pagamento", "Último pagamento")]
    elif tipo in ("parcelas", "recebimentos"):
        cond, params = filtro_cliente(q, carteira)
        campo_data = "p.pago_em" if tipo == "recebimentos" else "p.vencimento"
        params |= {"de": data_opcional(de, "De"), "ate": data_opcional(ate, "Até"),
                   "status": status if status in ("paga", "vencida", "a_vencer", "renegociada", "cancelada") else None}
        extra = "and p.status = 'paga'" if tipo == "recebimentos" else "and (%(status)s::text is null or p.status = %(status)s::text)"
        linhas = db.all(f"""
            select c.codigo, c.nome, c.cpf_cnpj, p.conta, p.nr_acordo, p.nr_parcela, p.vencimento, p.valor_previsto,
                   p.status, p.dias_atraso, p.valor_atualizado, p.pago_em, p.valor_pago, p.nosso_numero
            from virtua.v_parcelas p join virtua.v_clientes c on c.codigo = p.codigo_cliente
            where {cond} {extra}
              and (%(de)s::date is null or {campo_data} >= %(de)s::date)
              and (%(ate)s::date is null or {campo_data} <= %(ate)s::date)
            order by {campo_data}, c.nome""", params)
        cab = [("codigo", "Código"), ("nome", "Cliente"), ("cpf_cnpj", "CPF/CNPJ"), ("conta", "Contrato"), ("nr_acordo", "Acordo"),
               ("nr_parcela", "Parcela"), ("vencimento", "Vencimento"), ("valor_previsto", "Valor"), ("status", "Situação"),
               ("dias_atraso", "Dias de atraso"), ("valor_atualizado", "Com juros e multa"), ("pago_em", "Pago em"),
               ("valor_pago", "Valor pago"), ("nosso_numero", "Nosso número")]
    elif tipo == "inadimplencia":
        _, _, linhas = consulta_inadimplencia(q, carteira, faixa, de, ate, minimo, maximo, "valor", 100000, 0)
        cab = [("codigo", "Código"), ("nome", "Cliente"), ("cpf_cnpj", "CPF/CNPJ"), ("telefone", "Telefone"), ("email", "E-mail"),
               ("carteira", "Carteira"), ("parcelas_vencidas", "Parcelas vencidas"), ("valor_vencido", "Valor atrasado"),
               ("valor_atualizado", "Com juros e multa"), ("maior_atraso", "Maior atraso (dias)"),
               ("vencimento_mais_antigo", "Vencimento mais antigo"), ("ultimo_pagamento", "Último pagamento")]
    else:
        raise HTTPException(404, "Relatório desconhecido")
    return cab, linhas


ROTULOS_SITUACAO = {"ativa_em_dia": "Ativa · em dia", "ativa_em_atraso": "Ativa · em atraso",
                    "inadimplencia_antiga": "Inadimplência antiga", "paga": "Paga", "vencida": "Vencida",
                    "a_vencer": "A vencer", "renegociada": "Renegociada", "cancelada": "Cancelada"}


@app.get("/api/relatorios/{tipo}")
def relatorio(tipo: str, req: Request, formato: str = "xlsx", q: str = "", carteira: str = "", status: str = "",
              faixa: str = "", de: str = "", ate: str = "", minimo: str = "", maximo: str = ""):
    exigir(req)
    cab, linhas = dados_relatorio(tipo, q, carteira, status, faixa, de, ate, minimo, maximo)
    nome = f"{RELATORIOS[tipo]} {hoje().isoformat()}"

    def celula(chave, v):
        if chave in ("carteira", "status"):
            return ROTULOS_SITUACAO.get(v, v)
        return float(v) if isinstance(v, Decimal) else v

    if formato == "csv":
        import csv
        import io
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow([t for _, t in cab])
        def texto_csv(k, v):  # formato brasileiro, que o Excel abre direto
            if v is None:
                return ""
            if isinstance(v, Decimal):
                return f"{v:.2f}".replace(".", ",")
            if isinstance(v, date):
                return v.strftime("%d/%m/%Y")
            return celula(k, v)
        for l in linhas:
            w.writerow([texto_csv(k, l[k]) for k, _ in cab])
        corpo, mime, ext = ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "csv"
    else:
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = RELATORIOS[tipo][:31]
        ws.append([t for _, t in cab])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F5FBF")
        for l in linhas:
            ws.append([celula(k, l[k]) for k, _ in cab])
        for i, (k, _) in enumerate(cab, start=1):
            letra = ws.cell(1, i).column_letter
            if k in ("vencimento", "pago_em", "ultimo_pagamento", "vencimento_mais_antigo"):
                fmt, larg = "DD/MM/YYYY", 12
            elif k.startswith(("valor", "total")):
                fmt, larg = '"R$" #,##0.00', 14
            else:
                fmt, larg = None, max(12, min(40, max((len(str(l[k] or "")) for l in linhas[:500]), default=10) + 2))
            ws.column_dimensions[letra].width = larg
            if fmt:
                for c in ws[letra][1:]:
                    c.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        corpo, mime, ext = buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    from urllib.parse import quote
    return Response(corpo, media_type=mime,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nome)}.{ext}"})


# ---------- site ----------
app.mount("/static", StaticFiles(directory=config.ROOT / "public"), name="static")


@app.get("/{caminho:path}")
def site(caminho: str):
    if caminho.startswith("api/"):
        raise HTTPException(404)
    return FileResponse(config.ROOT / "public" / "index.html", headers={"Cache-Control": "no-cache"})
